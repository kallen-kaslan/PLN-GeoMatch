"""Browser acceptance tests for Data Master imports and source provenance."""

import codecs
import io
import re
from pathlib import Path
from zipfile import ZipFile

import httpx
import pytest
from openpyxl import Workbook
from playwright.sync_api import Page, expect


@pytest.mark.parametrize(
    "encoding, prefix",
    [
        ("cp1252", b""),
        ("utf-16", b""),
        ("utf-16-be", codecs.BOM_UTF16_BE),
        ("utf-8", b""),
    ],
)
def test_csv_encoding_matches_preview_and_saved_data(
    page: Page, servers: tuple[str, str], encoding: str, prefix: bytes
) -> None:
    text = "IDPEL;NAMA;ALAMAT\n000123;Ìsmi – Café;Jalan François\n"
    page.goto("/master/dil")
    page.get_by_role("button", name="Unggah berkas", exact=True).click()
    page.locator("input[type=file]").set_input_files(
        {
            "name": "encoded.csv",
            "mimeType": "text/csv",
            "buffer": prefix + text.encode(encoding),
        }
    )
    form = page.get_by_role("region", name="Impor master DIL")
    expect(form.get_by_role("cell", name="Ìsmi – Café", exact=True)).to_be_visible()
    page.get_by_role("button", name="Simpan dan validasi master").click()
    expect(
        page.get_by_text("encoded.csv berhasil disimpan.", exact=True)
    ).to_be_visible()
    expect(page.get_by_role("row").filter(has_text="000123")).to_contain_text(
        "Ìsmi – Café"
    )
    records = httpx.get(f"{servers[1]}/master/dil/records").json()
    assert records["items"][0]["fields"]["nama"] == "Ìsmi – Café"


@pytest.mark.parametrize(
    "header_row, leading_value", [(2, None), (2, "   "), (8, "   "), (70, "   ")]
)
def test_excel_header_after_empty_rows_is_previewed_and_imported(
    page: Page, servers: tuple[str, str], header_row: int, leading_value: str | None
) -> None:
    book = Workbook()
    sheet = book.active
    sheet.title = "DIL_SALDO_MASK_UP_202606_53321"
    sheet.cell(1, 1, leading_value)
    for column, value in enumerate(["IDPEL", "NAMA", "LAT", "LON"], start=1):
        sheet.cell(header_row, column, value)
    for offset in range(8):
        for column, value in enumerate(
            [f"00012345678{offset}", f"Pelanggan {offset}", "-6.2", "106.8"], start=1
        ):
            sheet.cell(header_row + 1 + offset, column, value)
    content = io.BytesIO()
    book.save(content)
    page.goto("/master/dil")
    page.get_by_role("button", name="Unggah berkas", exact=True).click()
    page.locator("input[type=file]").set_input_files(
        {
            "name": "dil-offset.xlsx",
            "mimeType": "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            "buffer": content.getvalue(),
        }
    )
    expect(page.get_by_role("alert")).to_have_count(0)
    expect(page.get_by_label("Kolom IDPEL", exact=True)).to_have_value("IDPEL")
    preview = page.get_by_role("region", name="Impor master DIL").get_by_role("table")
    expect(preview.get_by_role("row")).to_have_count(6)
    expect(preview.get_by_role("row").nth(1).get_by_role("cell").first).to_have_text(
        str(header_row + 1)
    )
    page.get_by_role("button", name="Simpan dan validasi master").click()
    expect(
        page.get_by_text("dil-offset.xlsx berhasil disimpan.", exact=True)
    ).to_be_visible()
    records = httpx.get(f"{servers[1]}/master/dil/records").json()
    assert records["total"] == 8
    assert [record["row_number"] for record in records["items"]] == list(
        range(header_row + 1, header_row + 9)
    )


@pytest.mark.parametrize("hidden", [False, True])
def test_excel_empty_first_sheet_uses_data_sheet(page: Page, hidden: bool) -> None:
    book = Workbook()
    book.active.title = "Kosong"
    if hidden:
        book.active.sheet_state = "hidden"
    data = book.create_sheet("DIL")
    data.append(["IDPEL", "NAMA", "KOORDINAT_Y", "KOORDINAT_X"])
    data.append(["000123456789", "Pelanggan DIL", "-6.2", "106.8"])
    content = io.BytesIO()
    book.save(content)
    page.goto("/master/dil")
    page.get_by_role("button", name="Unggah berkas", exact=True).click()
    page.locator("input[type=file]").set_input_files(
        {
            "name": "dil.xlsx",
            "mimeType": "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            "buffer": content.getvalue(),
        }
    )
    expect(page.get_by_label("Sheet yang diimpor")).to_have_value("DIL")
    expect(page.get_by_label("Kolom IDPEL", exact=True)).to_have_value("IDPEL")
    expect(page.get_by_role("alert")).to_have_count(0)
    if not hidden:
        page.get_by_label("Sheet yang diimpor").select_option("Kosong")
        expect(page.get_by_role("alert")).to_contain_text("Sheet")
        expect(
            page.get_by_role("button", name="Simpan dan validasi master")
        ).to_be_disabled()
        page.get_by_label("Sheet yang diimpor").select_option("DIL")
    page.get_by_role("button", name="Simpan dan validasi master").click()
    expect(page.get_by_text("dil.xlsx berhasil disimpan.", exact=True)).to_be_visible()
    expect(page.get_by_role("row").filter(has_text="000123456789")).to_contain_text(
        "Sheet: DIL"
    )


@pytest.mark.parametrize("header_row", [1, 70])
def test_excel_bad_dimensions_preview_matches_import_and_empty_sheet_can_recover(
    page: Page, servers: tuple[str, str], header_row: int
) -> None:
    book = Workbook()
    sheet = book.active
    sheet.title = "DIL_SALDO_MASK_UP_202606_53311"
    book.create_sheet("Kosong")
    for column, value in enumerate(["IDPEL", "NAMA", "LAT", "LON"], start=1):
        sheet.cell(header_row, column, value)
    for offset in range(8):
        for column, value in enumerate(
            [f"00012345678{offset}", f"Pelanggan {offset}", "-6.2", "106.8"], start=1
        ):
            sheet.cell(header_row + offset + 1, column, value)
    original = io.BytesIO()
    book.save(original)
    patched = io.BytesIO()
    with (
        ZipFile(io.BytesIO(original.getvalue())) as source,
        ZipFile(patched, "w") as target,
    ):
        for item in source.infolist():
            data = source.read(item.filename)
            if item.filename == "xl/worksheets/sheet1.xml":
                data = re.sub(
                    rb'<dimension ref="[^"]+"', b'<dimension ref="A1:A1"', data
                )
            target.writestr(item, data)
    requests = []
    page.on("request", lambda request: requests.append(request.url))
    page.goto("/master/dil")
    page.get_by_role("button", name="Unggah berkas", exact=True).click()
    page.locator("input[type=file]").set_input_files(
        {
            "name": "dil-dimension.xlsx",
            "mimeType": "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            "buffer": patched.getvalue(),
        }
    )
    expect(page.get_by_label("Kolom IDPEL", exact=True)).to_have_value(
        "IDPEL", timeout=20000
    )
    expect(page.get_by_label("Kolom Nama pelanggan", exact=True)).to_have_value("NAMA")
    form = page.get_by_role("region", name="Impor master DIL")
    preview = form.get_by_role("table")
    expect(preview.get_by_role("row").nth(1).get_by_role("cell").first).to_have_text(
        str(header_row + 1)
    )
    expect(preview.get_by_role("row")).to_have_count(6)
    assert httpx.get(f"{servers[1]}/master/dil/batches").json() == []
    if header_row > 64:
        assert any(url.endswith("/upload/preview") for url in requests)
    page.get_by_label("Sheet yang diimpor").select_option("Kosong")
    expect(page.get_by_role("alert")).to_contain_text("Kosong")
    expect(
        page.get_by_role("button", name="Simpan dan validasi master")
    ).to_be_disabled()
    page.get_by_label("Sheet yang diimpor").select_option(sheet.title)
    expect(page.get_by_label("Kolom IDPEL", exact=True)).to_have_value(
        "IDPEL", timeout=20000
    )
    expect(page.get_by_role("alert")).to_have_count(0)
    page.get_by_role("button", name="Simpan dan validasi master").click()
    expect(
        page.get_by_text("dil-dimension.xlsx berhasil disimpan.", exact=True)
    ).to_be_visible(timeout=20000)
    rows = httpx.get(f"{servers[1]}/master/dil/records").json()
    assert rows["total"] == 8
    assert [row["fields"]["idpel"] for row in rows["items"]] == [
        f"00012345678{i}" for i in range(8)
    ]
    assert [row["row_number"] for row in rows["items"]] == list(
        range(header_row + 1, header_row + 9)
    )


def test_excel_empty_workbook_can_select_sheet_and_replace_file(page: Page) -> None:
    book = Workbook()
    book.active.title = "Kosong"
    book.create_sheet("Kosong kedua")
    content = io.BytesIO()
    book.save(content)
    page.goto("/master/dil")
    page.get_by_role("button", name="Unggah berkas", exact=True).click()
    file_input = page.locator("input[type=file]")
    file_input.set_input_files(
        {
            "name": "empty.xlsx",
            "mimeType": "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            "buffer": content.getvalue(),
        }
    )
    expect(page.get_by_label("Sheet yang diimpor")).to_be_visible()
    expect(page.get_by_role("alert")).to_contain_text("Sheet")
    page.get_by_label("Sheet yang diimpor").select_option("Kosong kedua")
    expect(page.get_by_role("alert")).to_contain_text("Kosong kedua")
    file_input.set_input_files(
        {
            "name": "valid.csv",
            "mimeType": "text/csv",
            "buffer": b"IDPEL,NAMA\n001,Valid\n",
        }
    )
    expect(page.get_by_role("alert")).to_have_count(0)
    expect(page.get_by_label("Kolom IDPEL", exact=True)).to_have_value("IDPEL")


def test_csv_upload_mapping_details_and_category_isolation(
    page: Page, tmp_path: Path
) -> None:
    page.goto("/master/dil")
    expect(
        page.get_by_role("heading", name="Data Master DIL", exact=True)
    ).to_be_visible()
    page.get_by_role("button", name="Unggah berkas", exact=True).click()
    page.locator("input[type=file]").set_input_files(
        {
            "name": "dil-custom.csv",
            "mimeType": "text/csv",
            "buffer": b"CUSTOMER,CUSTOM_NAME,CODE_RAW,GARDU_CODE,Y_POS,X_POS\n000123456789,Budi,DL-01,G-03,-6.2,106.8\n000123456790,Siti,DL-02,G-04,,\n",
        }
    )
    page.get_by_label("Kolom IDPEL", exact=True).select_option("CUSTOMER")
    page.get_by_label("Kolom Nama pelanggan", exact=True).select_option("CUSTOM_NAME")
    page.get_by_label("Kolom Kode DIL", exact=True).select_option("CODE_RAW")
    page.get_by_label("Kolom Kode gardu", exact=True).select_option("GARDU_CODE")
    page.get_by_label("Kolom Latitude / koordinat Y", exact=True).select_option("Y_POS")
    page.get_by_label("Kolom Longitude / koordinat X", exact=True).select_option(
        "X_POS"
    )
    page.get_by_role("button", name="Simpan dan validasi master").click()
    expect(
        page.get_by_text("dil-custom.csv berhasil disimpan.", exact=True)
    ).to_be_visible()
    first_row = page.get_by_role("row").filter(has_text="000123456789")
    expect(first_row).to_contain_text("DL-01")
    expect(first_row).to_contain_text("G-03")
    first_row.get_by_role("button", name="Lihat detail baris 2").click()
    dialog = page.get_by_role("dialog")
    expect(dialog.get_by_text("Data asli berkas", exact=True)).to_be_visible()
    expect(dialog.get_by_text("CODE_RAW", exact=True)).to_be_visible()
    expect(dialog.get_by_text("000123456789", exact=True).first).to_be_visible()
    page.keyboard.press("Escape")
    expect(dialog).to_have_count(0)
    page.get_by_label("Status kelengkapan data").select_option("incomplete")
    expect(page.get_by_role("row").filter(has_text="000123456790")).to_be_visible()
    expect(page.get_by_role("row").filter(has_text="000123456789")).to_have_count(0)
    page.screenshot(path=str(tmp_path / "master-dil-desktop.png"), full_page=True)
    page.get_by_role("navigation", name="Jenis data master").get_by_role(
        "link", name="Prabayar", exact=True
    ).click()
    expect(
        page.get_by_role("heading", name="Data Master Prabayar", exact=True)
    ).to_be_visible()
    expect(page.get_by_role("row").filter(has_text="000123456790")).to_have_count(0)


def test_excel_sheet_selection_preserves_source_details(page: Page) -> None:
    book = Workbook()
    book.active.title = "Petunjuk"
    book.active.append(["INFO"])
    book.active.append(["Pilih sheet Data DIL"])
    sheet = book.create_sheet("Data DIL")
    sheet.append(
        [
            "NO",
            "V_BULAN_REKAP",
            "",
            "IDPEL",
            "NAMA",
            "NOMOR_GARDU",
            "KODEDLPD",
            "KDDK",
            "STATUS_DIL",
            "KOORDINAT_X",
            "KOORDINAT_Y",
            "NOMOR_METER_KWH",
            "NOMOR_METER_PREPAID",
        ]
    )
    sheet.append(
        [
            "7",
            "202609",
            "",
            "000123456789",
            "Sumber Excel",
            "G-99",
            "DL-99",
            "0007",
            "AKTIF",
            "106.8",
            "-6.2",
            "000123",
            "000456",
        ]
    )
    output = io.BytesIO()
    book.save(output)
    page.goto("/master/dil")
    page.get_by_role("button", name="Unggah berkas", exact=True).click()
    page.locator("input[type=file]").set_input_files(
        {
            "name": "reference.xlsx",
            "mimeType": "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            "buffer": output.getvalue(),
        }
    )
    page.get_by_label("Sheet yang diimpor").select_option("Data DIL")
    expect(page.get_by_label("Kolom IDPEL", exact=True)).to_have_value("IDPEL")
    expect(page.get_by_label("Kolom Tanggal baca / periode", exact=True)).to_have_value(
        "V_BULAN_REKAP"
    )
    expect(page.get_by_label("Kolom Kode gardu", exact=True)).to_have_value(
        "NOMOR_GARDU"
    )
    expect(page.get_by_label("Kolom KODEDLPD", exact=True)).to_have_value("KODEDLPD")
    expect(page.get_by_label("Kolom Kode DIL", exact=True)).to_have_value("")
    page.get_by_label("Periode berkas (opsional)").fill("2026-10")
    page.get_by_role("button", name="Simpan dan validasi master").click()
    expect(
        page.get_by_text("reference.xlsx berhasil disimpan.", exact=True)
    ).to_be_visible()
    row = page.get_by_role("row").filter(has_text="000123456789")
    expect(row).to_contain_text("Sheet: Data DIL")
    expect(row).to_contain_text("202609")
    expect(row).to_contain_text("G-99")
    expect(row).to_contain_text("KDDK: 0007")
    expect(row).to_contain_text("STATUS_DIL: AKTIF")
    row.get_by_role("button", name="Lihat detail baris 2").click()
    dialog = page.get_by_role("dialog")
    expect(dialog.get_by_text("Kolom_3", exact=True)).to_be_visible()
    expect(dialog.get_by_text("000123", exact=True).first).to_be_visible()
    expect(dialog.get_by_text("000456", exact=True).first).to_be_visible()


def test_pending_mapping_can_resume_without_duplicate_upload(
    page: Page, servers: tuple[str, str]
) -> None:
    response = httpx.post(
        f"{servers[1]}/upload/",
        data={"dataset_type": "dil"},
        files={"file": ("pending.csv", b"C\n000123456789\n", "text/csv")},
    )
    response.raise_for_status()
    source_id = response.json()["id"]
    page.goto("/master/dil")
    page.get_by_label("Berkas sumber", exact=True).select_option(str(source_id))
    page.get_by_role("button", name="Lanjutkan mapping", exact=True).click()
    page.get_by_label("Kolom IDPEL", exact=True).select_option("C")
    page.get_by_role("button", name="Simpan dan validasi master").click()
    expect(
        page.get_by_text("pending.csv berhasil disimpan.", exact=True)
    ).to_be_visible()
    assert len(httpx.get(f"{servers[1]}/master/dil/batches").json()) == 1


def test_mobile_navigation_and_upload_separation(page: Page) -> None:
    page.set_viewport_size({"width": 390, "height": 844})
    page.goto("/master/pascabayar")
    expect(
        page.get_by_role("heading", name="Data Master Pascabayar", exact=True)
    ).to_be_visible()
    assert page.evaluate("document.documentElement.scrollWidth <= window.innerWidth")
    page.get_by_role("button", name="Buka menu navigasi", exact=True).click()
    drawer = page.get_by_role("dialog", name="Menu Navigasi Mobile")
    drawer.get_by_role("link", name="DIL", exact=True).click()
    expect(
        page.get_by_role("heading", name="Data Master DIL", exact=True)
    ).to_be_visible()
    page.goto("/upload")
    expect(page.get_by_role("heading", name="Data Koreksi", exact=True)).to_be_visible()
    expect(
        page.get_by_role("button", name="1. Data Pascabayar", exact=False)
    ).to_have_count(0)
