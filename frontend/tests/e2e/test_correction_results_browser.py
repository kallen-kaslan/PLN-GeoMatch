"""Map choices appear immediately in results, XLSX, and process history."""

from pathlib import Path

import httpx
import pytest
from openpyxl import load_workbook
from playwright.sync_api import Page, expect

from test_corrections_browser import prepare, upload


@pytest.mark.parametrize("kind", ["idpel_kosong", "idpel_tidak_ketemu"])
def test_partial_results_export_history_resume_replace_and_cancel(
    page: Page, servers: tuple[str, str], tmp_path: Path, kind: str
) -> None:
    api = servers[1]
    review = prepare(api, candidates=2)
    batch = upload(api, kind, "IDPEL,NAMA,LAT,LON\n" + "".join(f"{700+i if kind == 'idpel_tidak_ketemu' else ''},Target {i},-6.2,106.8\n" for i in range(5)))
    response = httpx.post(f"{api}/corrections/{kind}/checks", json={"batch_id": batch, "pdpj_review_id": review, "radius_meters": 10, "actor": "Petugas", "request_key": "results_browser_1", "incremental": True})
    response.raise_for_status()
    check = response.json()
    process = check["process_id"]
    base = f"{api}/corrections/{kind}/checks/{process}"
    run = check["latest_run"]
    for _ in range(2):
        response = httpx.post(base + "/advance", json={"run_id": run["id"], "revision": run["processing"]["revision"], "chunk_size": 2})
        response.raise_for_status()
        run = response.json()
    httpx.post(base + "/pause", json={"run_id": run["id"]}).raise_for_status()
    page.goto(f"/peta-koreksi?category={kind}&batch_id={batch}&process_id={process}")
    page.get_by_role("button", name="Pilih target baris 2", exact=True).click()
    page.get_by_label("Petugas peta", exact=True).fill("Petugas Hasil")
    page.get_by_role("button", name="Tampilkan kandidat target", exact=True).click()
    candidates = page.get_by_role("region", name="Kandidat target peta", exact=True)
    candidates.get_by_role("button", name="Pilih kandidat 1", exact=True).click()
    expect(page.get_by_role("region", name="Detail target terpilih", exact=True)).to_contain_text("Pilihan aktif: 002")
    page.get_by_role("link", name="Lihat Hasil Koreksi", exact=True).click()
    expect(page.get_by_role("heading", name="Hasil Koreksi", exact=True)).to_be_visible()
    expect(page.get_by_text("Hasil sebagian: 2 dari 5 baris", exact=False)).to_be_visible()
    results = page.get_by_role("region", name="Daftar hasil koreksi", exact=True)
    expect(results).to_contain_text("002")
    expect(results).to_contain_text("Petugas Hasil")
    expect(results.get_by_role("row")).to_have_count(2)
    with page.expect_download() as download:
        page.get_by_role("button", name="Ekspor hasil XLSX", exact=True).click()
    assert download.value.suggested_filename.endswith(".xlsx")
    destination = tmp_path / "result.xlsx"
    download.value.save_as(destination)
    book = load_workbook(destination)
    rows = list(book["Hasil Koreksi"].values)
    assert len(rows) == 2
    exported = dict(zip(rows[0], rows[1]))
    assert exported["idpelanggan"] == "002"
    assert exported["latitudey"] == "-6.2"
    assert exported["x_acc"] == "106.8"
    assert exported["y_acc"] == "-6.2"
    assert "dil.csv" in exported["sumber file (sheet)"]
    assert book["Hasil Koreksi"].freeze_panes is None
    assert book["Hasil Koreksi"]["A1"].fill.patternType is None
    assert book["Hasil Koreksi"]["B2"].data_type == "s"
    book.close()
    page.reload()
    expect(results).to_contain_text("002")
    page.get_by_role("button", name="Riwayat hasil baris 2", exact=True).click()
    audit = page.get_by_role("region", name="Riwayat keputusan koreksi", exact=True)
    expect(audit).to_contain_text("Revisi 1")
    page.get_by_role("link", name="Buka Riwayat Proses", exact=True).click()
    expect(page.get_by_role("table")).to_contain_text("2 baris tersimpan")
    expect(page.get_by_role("table")).to_contain_text("3 belum diproses")
    expect(page.get_by_role("table")).to_contain_text("1 pilihan aktif")
    expect(audit).to_contain_text("Kandidat dipilih: 002")
    with page.expect_download() as download:
        page.get_by_role("button", name=f"Unduh hasil XLSX proses {process}", exact=True).last.click()
    download.value.save_as(tmp_path / "history-results.xlsx")
    book = load_workbook(tmp_path / "history-results.xlsx")
    assert book["Hasil Koreksi"]["B2"].value == "002"
    book.close()
    with page.expect_download() as download:
        audit.get_by_role("button", name="Unduh riwayat XLSX", exact=True).click()
    download.value.save_as(tmp_path / "history.xlsx")
    book = load_workbook(tmp_path / "history.xlsx")
    assert book["Riwayat Pilihan"]["B2"].value == "002"
    book.close()
    page.get_by_role("link", name="Hasil", exact=True).click()
    response = httpx.post(base + "/resume", json={"run_id": run["id"]})
    response.raise_for_status()
    run = response.json()
    for _ in range(12):
        if run["status"] == "completed":
            break
        response = httpx.post(base + "/advance", json={"run_id": run["id"], "revision": run["processing"]["revision"], "chunk_size": 2})
        response.raise_for_status()
        run = response.json()
    assert run["status"] == "completed"
    page.get_by_role("button", name="Muat ulang hasil", exact=True).click()
    expect(results).to_contain_text("002")
    page.get_by_role("link", name="Buka Peta Koreksi", exact=True).click()
    page.get_by_role("button", name="Pilih target baris 2", exact=True).click()
    page.get_by_label("Petugas peta", exact=True).fill("Petugas Hasil")
    page.get_by_role("button", name="Tampilkan kandidat target", exact=True).click()
    page.get_by_label("Alasan perubahan pilihan", exact=True).fill("Meter kandidat kedua lebih sesuai")
    candidates.get_by_role("button", name="Pilih kandidat 2", exact=True).click()
    expect(page.get_by_role("region", name="Detail target terpilih", exact=True)).to_contain_text("Pilihan aktif: 003")
    page.get_by_role("link", name="Lihat Hasil Koreksi", exact=True).click()
    expect(results).to_contain_text("003")
    expect(results.get_by_role("row")).to_have_count(2)
    expect(audit).to_contain_text("Revisi 2")
    expect(audit).to_contain_text("Kandidat dipilih: 002")
    expect(audit).to_contain_text("Kandidat dipilih: 003")
    page.get_by_role("link", name="Buka Peta Koreksi", exact=True).click()
    page.get_by_role("button", name="Pilih target baris 2", exact=True).click()
    page.get_by_label("Petugas peta", exact=True).fill("Petugas Hasil")
    page.get_by_label("Alasan perubahan pilihan", exact=True).fill("Perlu verifikasi lapangan")
    page.get_by_role("button", name="Batalkan pilihan", exact=True).click()
    expect(page.get_by_text("Pilihan dibatalkan. Riwayat tetap tersimpan.", exact=True)).to_be_visible()
    page.get_by_role("link", name="Lihat Hasil Koreksi", exact=True).click()
    expect(results).to_contain_text("Tidak ada hasil sesuai filter.")
    page.get_by_label("Filter hasil pilihan", exact=True).select_option("cancelled")
    expect(results.get_by_role("row")).to_have_count(2)
    expect(results).to_contain_text("Pilihan dibatalkan")
    expect(audit).to_contain_text("Revisi 3")
    expect(audit).to_contain_text("Perlu verifikasi lapangan")
    with page.expect_download() as download:
        audit.get_by_role("button", name="Unduh riwayat XLSX", exact=True).click()
    download.value.save_as(tmp_path / "history-cancelled.xlsx")
    book = load_workbook(tmp_path / "history-cancelled.xlsx")
    rows = list(book["Riwayat Pilihan"].values)
    assert len(rows) == 4
    events = [dict(zip(rows[0], row)) for row in rows[1:]]
    assert [event["revisi_pilihan"] for event in events] == [3, 2, 1]
    assert events[0]["idpelanggan"] is None
    assert events[0]["status"] == "Pilihan dibatalkan"
    assert events[0]["x_acc"] is None and events[0]["y_acc"] is None
    assert events[1]["x_acc"] == "106.8" and events[1]["y_acc"] == "-6.2"
    assert events[1]["idpelanggan"] == "003"
    assert events[2]["idpelanggan"] == "002"
    book.close()
    page.set_viewport_size({"width": 390, "height": 844})
    expect(page.get_by_role("heading", name="Hasil Koreksi", exact=True)).to_be_visible()
    assert page.evaluate("document.documentElement.scrollWidth <= window.innerWidth")
