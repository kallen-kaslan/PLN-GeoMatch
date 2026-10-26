"""Reuse original workbooks across reloads without activating unconfirmed sheets."""

from io import BytesIO

import httpx
from openpyxl import Workbook
from playwright.sync_api import Page, expect


def workbook() -> bytes:
    book = Workbook()
    book.active.title = "Januari"
    book.active.append(["IDPEL", "NAMA", "LAT", "LON"])
    book.active.append(["001", "Pelanggan Januari", -6.2, 106.8])
    other = book.create_sheet("Februari")
    other.append(["IDPEL", "NAMA", "LAT", "LON"])
    other.append(["002", "Pelanggan Februari", -6.3, 106.9])
    book.create_sheet("Kosong")
    stream = BytesIO()
    book.save(stream)
    return stream.getvalue()


def open_import(page: Page) -> None:
    page.get_by_role("button", name="Unggah berkas", exact=True).click()


def choose_original(page: Page, contents: bytes) -> None:
    page.locator("input[type=file]").set_input_files(
        {
            "name": "riwayat.xlsx",
            "mimeType": "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            "buffer": contents,
        }
    )
    expect(page.get_by_label("Sheet yang diimpor", exact=True)).to_be_enabled(
        timeout=20000
    )


def test_sheet_history_checks_survive_reload_and_other_sheet_uses_same_original(
    page: Page, servers: tuple[str, str]
) -> None:
    uploads = []
    page.on(
        "request",
        lambda request: (
            uploads.append(request.url)
            if request.method == "POST" and request.url.endswith("/upload/files")
            else None
        ),
    )
    page.goto("/master/dil")
    open_import(page)
    choose_original(page, workbook())
    save = page.get_by_role("button", name="Simpan dan validasi master", exact=True)
    expect(save).to_be_enabled()
    assert httpx.get(servers[1] + "/master/dil/summary").json()["total_rows"] == 0
    assert httpx.get(servers[1] + "/upload/batches").json() == []
    save.click()
    expect(
        page.get_by_text("riwayat.xlsx berhasil disimpan.", exact=True)
    ).to_be_visible(timeout=20000)
    page.reload()
    open_import(page)
    page.get_by_label("Riwayat berkas unggahan", exact=True).select_option(index=1)
    selector = page.get_by_label("Sheet yang diimpor", exact=True)
    expect(selector).to_be_enabled()
    expect(selector.locator("option[value=Januari]")).to_have_text(
        "✓ Januari"
    )
    expect(save).to_be_disabled()
    selector.select_option("Februari")
    expect(save).to_be_enabled()
    expect(page.get_by_text("Pelanggan Februari", exact=True)).to_be_visible()
    assert httpx.get(servers[1] + "/master/dil/summary").json()["total_rows"] == 1
    # Even leaving/reopening the preview must not silently add it to reference data.
    page.reload()
    open_import(page)
    page.get_by_label("Riwayat berkas unggahan", exact=True).select_option(index=1)
    page.get_by_label("Sheet yang diimpor", exact=True).select_option("Februari")
    expect(save).to_be_enabled()
    save.click()
    expect(
        page.get_by_text("riwayat.xlsx berhasil disimpan.", exact=True)
    ).to_be_visible(timeout=20000)
    assert httpx.get(servers[1] + "/master/dil/summary").json()["total_rows"] == 2
    assert len(uploads) == 1
    open_import(page)
    page.get_by_label("Riwayat berkas unggahan", exact=True).select_option(index=1)
    selector = page.get_by_label("Sheet yang diimpor", exact=True)
    expect(selector.locator("option[value=Januari]")).to_have_text(
        "✓ Januari"
    )
    expect(selector.locator("option[value=Februari]")).to_have_text(
        "✓ Februari"
    )
    selector.select_option("Februari")
    expect(save).to_be_disabled()
    assert len(httpx.get(servers[1] + "/upload/batches").json()) == 2


def test_preview_empty_sheet_recovery_and_checkmarks_follow_period(
    page: Page, servers: tuple[str, str]
) -> None:
    page.goto("/master/dil")
    open_import(page)
    choose_original(page, workbook())
    page.locator("input[type=month]").fill("2026-01")
    save = page.get_by_role("button", name="Simpan dan validasi master", exact=True)
    save.click()
    expect(
        page.get_by_text("riwayat.xlsx berhasil disimpan.", exact=True)
    ).to_be_visible(timeout=20000)
    open_import(page)
    page.get_by_label("Riwayat berkas unggahan", exact=True).select_option(index=1)
    selector = page.get_by_label("Sheet yang diimpor", exact=True)
    expect(selector).to_be_enabled()
    expect(selector.locator("option[value=Januari]")).to_have_text(
        "Januari"
    )
    page.locator("input[type=month]").fill("2026-01")
    expect(selector.locator("option[value=Januari]")).to_have_text(
        "✓ Januari"
    )
    expect(save).to_be_disabled()
    selector.select_option("Kosong")
    expect(page.get_by_role("alert")).to_contain_text("tidak memiliki nama kolom")
    expect(save).to_be_disabled()
    selector.select_option("Februari")
    expect(save).to_be_enabled()
    page.get_by_role("button", name="Tutup formulir unggah", exact=True).click()
    assert httpx.get(servers[1] + "/master/dil/summary").json()["total_rows"] == 1


def test_same_file_history_can_be_reused_in_pdpj_without_master_checkmarks(
    page: Page, servers: tuple[str, str]
) -> None:
    page.goto("/master/dil")
    open_import(page)
    choose_original(page, workbook())
    page.get_by_role("button", name="Simpan dan validasi master", exact=True).click()
    expect(
        page.get_by_text("riwayat.xlsx berhasil disimpan.", exact=True)
    ).to_be_visible(timeout=20000)
    page.goto("/pdpj")
    page.get_by_role("button", name="Unggah PDPJ", exact=True).click()
    page.get_by_label("Riwayat berkas unggahan", exact=True).select_option(index=1)
    selector = page.get_by_label("Sheet yang diimpor", exact=True)
    expect(selector).to_be_enabled()
    expect(selector.locator("option[value=Januari]")).to_have_text(
        "Januari"
    )
    save = page.get_by_role("button", name="Simpan dan validasi PDPJ", exact=True)
    expect(save).to_be_enabled()
    assert len(httpx.get(servers[1] + "/upload/batches").json()) == 1
    save.click()
    expect(page.get_by_text("riwayat.xlsx tersimpan:", exact=False)).to_be_visible(
        timeout=20000
    )
    assert len(httpx.get(servers[1] + "/upload/files").json()) == 1
