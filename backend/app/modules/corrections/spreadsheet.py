"""Build plain XLSX tables while preserving identifier and coordinate text."""

from collections.abc import Iterator, Sequence
from tempfile import SpooledTemporaryFile
from typing import Any, BinaryIO, cast

from fastapi.responses import StreamingResponse
from openpyxl import Workbook
from openpyxl.cell import WriteOnlyCell
from openpyxl.utils import get_column_letter
from openpyxl.worksheet._write_only import WriteOnlyWorksheet

XLSX_MEDIA_TYPE = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
RESULT_HEADERS = [
    "noagenda",
    "idpelanggan",
    "verif",
    "jenispelanggan",
    "latitudey",
    "longitudex",
    "sumber file (sheet)",
    "sel (baris ke berapa)",
    "x_acc",
    "y_acc",
]
RESULT_WIDTHS = [26, 22, 16, 22, 24, 24, 58, 26, 24, 24]


class Spreadsheet:
    """Write rows incrementally; never interpret uploaded strings as formulas."""

    def __init__(self) -> None:
        self.book = Workbook(write_only=True)
        self.book.properties.creator = "PLN GeoMatch"

    def sheet(
        self, title: str, headers: Sequence[str], widths: Sequence[int]
    ) -> WriteOnlyWorksheet:
        sheet = self.book.create_sheet(title)
        for index, width in enumerate(widths, 1):
            sheet.column_dimensions[get_column_letter(index)].width = width
        sheet.append(list(headers))
        return sheet

    def append(self, sheet: WriteOnlyWorksheet, values: Sequence[Any]) -> None:
        cells = []
        for value in values:
            cell = WriteOnlyCell(sheet, value=value)
            if isinstance(value, str):
                # XLSX strings preserve leading zeroes, decimal digits, and literal '='.
                cell.data_type = "s"
                cell.number_format = "@"
            cells.append(cell)
        sheet.append(cells)

    def response(self, filename: str) -> StreamingResponse:
        output = cast(
            BinaryIO,
            SpooledTemporaryFile(max_size=8 * 1024 * 1024, mode="w+b"),  # noqa: SIM115 -- response closes it
        )
        try:
            self.book.save(output)
            output.seek(0)
        except BaseException:
            output.close()
            raise

        def content() -> Iterator[bytes]:
            try:
                while chunk := output.read(64 * 1024):
                    yield chunk
            finally:
                output.close()

        return StreamingResponse(
            content(),
            media_type=XLSX_MEDIA_TYPE,
            headers={"Content-Disposition": f'attachment; filename="{filename}"'},
        )
