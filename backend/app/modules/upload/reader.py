"""Read uploads from their spooled file without materializing the whole dataset."""

import codecs
import csv
from collections.abc import Iterator
from contextlib import contextmanager
from typing import BinaryIO

import pandas as pd
from fastapi import HTTPException
from openpyxl import load_workbook

from .normalization import clean_text, normalize_source_headers

MAX_UPLOAD_BYTES = 500 * 1024 * 1024
IMPORT_CHUNK_ROWS = 1000
ENCODING_CHUNK_BYTES = 1024 * 1024


def validate_text_encoding(source: BinaryIO, encoding: str) -> None:
    """Check the entire stream, including characters beyond the preview sample."""
    decoder = codecs.getincrementaldecoder(encoding)(errors="strict")
    source.seek(0)
    try:
        while chunk := source.read(ENCODING_CHUNK_BYTES):
            decoder.decode(chunk, final=False)
        decoder.decode(b"", final=True)
    finally:
        source.seek(0)


def detect_csv_encoding(source: BinaryIO) -> str:
    source.seek(0)
    marker = source.read(4)
    source.seek(0)
    if marker.startswith((codecs.BOM_UTF32_LE, codecs.BOM_UTF32_BE)):
        raise ValueError("CSV UTF-32 belum didukung. Simpan ulang sebagai CSV UTF-8.")
    declared = (
        "utf-8-sig"
        if marker.startswith(codecs.BOM_UTF8)
        else "utf-16"
        if marker.startswith((codecs.BOM_UTF16_LE, codecs.BOM_UTF16_BE))
        else None
    )
    try:
        if declared:
            validate_text_encoding(source, declared)
            return declared
        try:
            validate_text_encoding(source, "utf-8-sig")
            return "utf-8-sig"
        except UnicodeDecodeError:
            validate_text_encoding(source, "cp1252")
            return "cp1252"
    except UnicodeError as exc:
        raise ValueError(
            "Encoding CSV tidak dapat dibaca. Simpan ulang sebagai CSV UTF-8 atau UTF-16."
        ) from exc


def check_file_size(source: BinaryIO) -> None:
    source.seek(0, 2)
    size = source.tell()
    source.seek(0)
    if size > MAX_UPLOAD_BYTES:
        raise HTTPException(status_code=413, detail="Ukuran file maksimal 500 MB.")


@contextmanager
def read_upload(source: BinaryIO, filename: str, sheet_name: str | None):
    if filename.lower().endswith(".csv"):
        encoding = detect_csv_encoding(source)
        # A sample may end inside a multibyte character; retain the partial codepoint.
        decoder = codecs.getincrementaldecoder(encoding)(errors="strict")
        sample = decoder.decode(source.read(32768), final=False)
        source.seek(0)
        try:
            separator = csv.Sniffer().sniff(sample, delimiters=",;\t|").delimiter
        except csv.Error:
            separator = ","
        with pd.read_csv(
            source,
            encoding=encoding,
            dtype=str,
            keep_default_na=False,
            skip_blank_lines=False,
            header=None,
            sep=separator,
            chunksize=IMPORT_CHUNK_ROWS,
        ) as reader:

            def csv_rows() -> Iterator[tuple[int, dict[str, str]]]:
                headers = None
                row_number = 2
                for frame in reader:
                    if headers is None:
                        if frame.empty:
                            raise ValueError("Berkas tidak memiliki header atau data")
                        headers = normalize_source_headers(frame.iloc[0].tolist())
                        frame = frame.iloc[1:]
                    frame.columns = headers
                    for row in frame.to_dict(orient="records"):
                        yield row_number, row
                        row_number += 1

            yield None, csv_rows()
        return

    if not filename.lower().endswith(".xlsx"):
        raise ValueError("Gunakan berkas CSV atau XLSX")
    workbook = load_workbook(source, read_only=True, data_only=True)
    try:
        selected_sheet = sheet_name or workbook.sheetnames[0]
        if selected_sheet not in workbook.sheetnames:
            raise ValueError(f"Sheet '{selected_sheet}' tidak ditemukan")
        worksheet = workbook[selected_sheet]
        # Some spreadsheet exporters write A1 as the dimension even when the
        # actual cells extend further. Stream the cells instead of trusting it.
        worksheet.reset_dimensions()

        def excel_rows() -> Iterator[tuple[int, dict[str, str]]]:
            rows = enumerate(worksheet.iter_rows(values_only=True), start=1)
            for _, values in rows:
                if any(clean_text(value) for value in values):
                    headers = normalize_source_headers(list(values))
                    break
            else:
                raise ValueError(f"Sheet '{selected_sheet}' tidak memiliki nama kolom")
            # Preserve internal blank rows, but ignore Excel's trailing empty rows.
            pending_blank = 0
            for row_number, values in rows:
                if all(value is None or value == "" for value in values):
                    pending_blank += 1
                    continue
                for blank_number in range(row_number - pending_blank, row_number):
                    yield blank_number, dict.fromkeys(headers, "")
                pending_blank = 0
                yield (
                    row_number,
                    {
                        header: ""
                        if index >= len(values) or values[index] is None
                        else str(values[index])
                        for index, header in enumerate(headers)
                    },
                )

        yield selected_sheet, excel_rows()
    finally:
        workbook.close()
