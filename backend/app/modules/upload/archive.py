"""Retain original files; previewing an archived sheet never creates reference data."""

import hashlib
import os
import re
from itertools import islice
from pathlib import Path
from typing import Annotated, Any
from uuid import uuid4
from zipfile import BadZipFile

from fastapi import APIRouter, Depends, File, HTTPException, Query, UploadFile
from openpyxl import load_workbook
from pydantic import BaseModel, ConfigDict, Field
from sqlalchemy.exc import IntegrityError, SQLAlchemyError
from sqlalchemy.orm import Session

from app.config import settings
from app.database import get_db
from app.models.upload import DatasetType, UploadBatch
from app.models.upload_file import UploadFileArchive, UploadSheetImport
from app.modules.pdpj.service import lock_review_session

from .ingestion import import_raw_sheet
from .reader import check_file_size, read_upload
from .storage import database_write_error

router = APIRouter()
DbSession = Annotated[Session, Depends(get_db)]


def archive_path(key: str) -> Path:
    if not re.fullmatch(r"[a-f0-9]{32}", key):
        raise ValueError("Kunci penyimpanan berkas tidak valid")
    root = Path(settings.UPLOAD_STORAGE_DIR).resolve()
    target = (root / key).resolve()
    if target.parent != root:
        raise ValueError("Lokasi berkas tidak valid")
    return target


def get_archive(db: Session, file_id: int) -> UploadFileArchive:
    archive = db.get(UploadFileArchive, file_id)
    if archive is None:
        raise HTTPException(404, "Berkas dalam riwayat tidak ditemukan")
    return archive


def archive_info(db: Session, archive: UploadFileArchive) -> dict[str, Any]:
    imports = (
        db.query(UploadSheetImport, UploadBatch)
        .join(UploadBatch, UploadBatch.id == UploadSheetImport.batch_id)
        .filter(UploadSheetImport.file_id == archive.id)
        .order_by(UploadBatch.id)
        .all()
    )
    return {
        "id": archive.id,
        "filename": archive.filename,
        "size_bytes": archive.size_bytes,
        "sheet_names": archive.sheet_names,
        "created_at": archive.created_at,
        "imports": [
            {
                "batch_id": batch.id,
                "dataset_type": batch.dataset_type,
                "source_sheet": batch.source_sheet,
                "period": batch.period,
                "status": batch.status,
                "total_rows": batch.total_rows,
            }
            for _, batch in imports
        ],
    }


@router.get("/files")
def list_files(db: DbSession) -> list[dict[str, Any]]:
    return [
        archive_info(db, item)
        for item in db.query(UploadFileArchive)
        .order_by(UploadFileArchive.id.desc())
        .all()
    ]


@router.post("/files")
def store_file(file: Annotated[UploadFile, File()], db: DbSession) -> dict[str, Any]:
    key = uuid4().hex
    path = archive_path(key)
    committed = False
    created = False
    try:
        filename = file.filename or ""
        if not filename.lower().endswith((".xlsx", ".csv")):
            raise ValueError("Gunakan berkas CSV atau XLSX")
        check_file_size(file.file)
        names: list[str] = []
        if filename.lower().endswith(".xlsx"):
            book = load_workbook(file.file, read_only=True, data_only=True)
            try:
                names = book.sheetnames
                if not names:
                    raise ValueError("Berkas Excel tidak memiliki sheet")
            finally:
                book.close()
        file.file.seek(0)
        path.parent.mkdir(parents=True, exist_ok=True)
        digest = hashlib.sha256()
        size = 0
        with path.open("xb") as destination:
            created = True
            while chunk := file.file.read(1024 * 1024):
                digest.update(chunk)
                size += len(chunk)
                destination.write(chunk)
        existing = (
            db.query(UploadFileArchive).filter_by(sha256=digest.hexdigest()).first()
        )
        if existing:
            original = archive_path(existing.storage_key)
            if not original.exists():
                os.replace(path, original)
            else:
                path.unlink()
            return archive_info(db, existing)
        archive = UploadFileArchive(
            filename=filename,
            sha256=digest.hexdigest(),
            storage_key=key,
            size_bytes=size,
            sheet_names=names,
        )
        db.add(archive)
        try:
            db.commit()
            committed = True
        except IntegrityError:
            db.rollback()
            existing = (
                db.query(UploadFileArchive).filter_by(sha256=digest.hexdigest()).first()
            )
            if existing is None:
                raise
            path.unlink(missing_ok=True)
            return archive_info(db, existing)
        return archive_info(db, archive)
    except SQLAlchemyError as exc:
        if created and not committed:
            path.unlink(missing_ok=True)
        raise database_write_error(db, exc, "riwayat berkas") from exc
    except (OSError, ValueError, ImportError, BadZipFile) as exc:
        if created and not committed:
            path.unlink(missing_ok=True)
        raise HTTPException(400, f"Berkas tidak dapat disimpan: {exc!s}") from exc
    finally:
        file.file.close()


class SheetRequest(BaseModel):
    model_config = ConfigDict(extra="forbid")
    dataset_type: DatasetType
    sheet_name: str | None = Field(None, max_length=128)
    period: str | None = Field(None, pattern=r"^\d{4}-(0[1-9]|1[0-2])$")


def selected_sheet(archive: UploadFileArchive, name: str | None) -> str | None:
    if not archive.sheet_names:
        if name is not None:
            raise HTTPException(422, "Berkas CSV tidak memiliki sheet")
        return None
    if name not in archive.sheet_names:
        raise HTTPException(422, "Pilih sheet yang tersedia dalam riwayat berkas")
    return name


@router.get("/files/{file_id}/preview")
def preview_sheet(
    file_id: int, db: DbSession, sheet_name: str | None = Query(None, max_length=128)
) -> dict[str, Any]:
    archive = get_archive(db, file_id)
    name = selected_sheet(archive, sheet_name)
    path = archive_path(archive.storage_key)
    if not path.is_file():
        raise HTTPException(
            410,
            "Berkas asli tidak tersedia. Unggah kembali berkas yang sama untuk memulihkannya.",
        )
    try:
        with (
            path.open("rb") as source,
            read_upload(source, archive.filename, name) as (sheet, rows),
        ):
            preview = list(islice(rows, 5))
            if not preview:
                raise ValueError("Sheet tidak memiliki baris data setelah nama kolom")
            return {
                "headers": list(preview[0][1]),
                "rows": [row for _, row in preview],
                "headerRow": preview[0][0] - 1,
                "source_sheet": sheet,
            }
    except (OSError, ValueError, ImportError, BadZipFile) as exc:
        raise HTTPException(400, f"Pratinjau tidak dapat dibaca: {exc!s}") from exc


@router.post("/files/{file_id}/import-sheet")
def import_sheet(file_id: int, req: SheetRequest, db: DbSession) -> dict[str, Any]:
    try:
        lock_review_session(db)
        archive = (
            db.query(UploadFileArchive)
            .filter_by(id=file_id)
            .populate_existing()
            .with_for_update()
            .first()
        )
        if archive is None:
            raise HTTPException(404, "Berkas dalam riwayat tidak ditemukan")
        name = selected_sheet(archive, req.sheet_name)
        existing = (
            db.query(UploadSheetImport)
            .filter_by(
                file_id=file_id,
                dataset_type=req.dataset_type.value,
                sheet_key=name or "",
                period_key=req.period or "",
            )
            .first()
        )
        if existing:
            batch = db.get(UploadBatch, existing.batch_id)
            assert batch is not None
            return {"id": batch.id, "status": batch.status, "filename": batch.filename}
        path = archive_path(archive.storage_key)
        if not path.is_file():
            raise HTTPException(
                410, "Berkas asli tidak tersedia. Unggah kembali berkas yang sama."
            )
        with path.open("rb") as source:
            batch = import_raw_sheet(
                db, source, archive.filename, req.dataset_type, name, req.period
            )
        db.add(
            UploadSheetImport(
                file_id=file_id,
                batch_id=batch.id,
                dataset_type=req.dataset_type.value,
                sheet_key=name or "",
                period_key=req.period or "",
            )
        )
        db.commit()
        return {"id": batch.id, "status": batch.status, "filename": batch.filename}
    except SQLAlchemyError as exc:
        raise database_write_error(db, exc, "sheet") from exc
    except (OSError, ValueError, ImportError, BadZipFile) as exc:
        db.rollback()
        raise HTTPException(400, f"Sheet tidak dapat disimpan: {exc!s}") from exc
