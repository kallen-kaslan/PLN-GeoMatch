"""All-source matching, immutable legacy checkpoints, and selected coordinate provenance."""

from io import BytesIO

import pytest
from fastapi.testclient import TestClient
from openpyxl import load_workbook
from sqlalchemy import create_engine
from sqlalchemy.orm import Session

from app.correction_observation_schema import upgrade_correction_observations
from app.models.correction import CorrectionCandidateEvidence, CorrectionTarget
from app.models.master_sisa import (
    CorrectionMasterObservation,
    CorrectionObservationCatalog,
)
from app.models.matching import MatchCandidate, MatchingRun
from app.modules.corrections import observations
from tests.test_correction_progress import advance, finish, path
from tests.test_corrections import upload


def scenario(
    client: TestClient, kind: str = "idpel_kosong", latest_lat: str = "-6.21"
) -> dict:
    upload(
        client,
        "pascabayar",
        "IDPEL,LAT,LON,TGLBACA\n002,-6.2002,106.8,2025-01-01\n003,-6.20005,106.8,2025-01-01\n",
    )
    upload(
        client,
        "pascabayar",
        f"IDPEL,LAT,LON,TGLBACA\n002,{latest_lat},106.8,2026-01-01\n003,-6.21,106.8,2026-01-01\n",
    )
    upload(
        client,
        "dil",
        "IDPEL,LAT,LON\n001,-6.2,106.8\n002,-6.2001,106.8\n003,-6.2003,106.8\n",
    )
    batch = upload(client, "pdpj", "IDPEL,LAT,LON\n001,-6.2,106.8\n")
    review = client.post("/pdpj/reviews", json={"batch_id": batch}).json()
    result = client.post(
        f"/pdpj/reviews/{review['id']}/complete",
        json={"revision": review["revision"], "actor": "Uji"},
    )
    assert result.status_code == 200
    targets = (
        "999,-6.2,106.8\n998,-6.2,106.8\n"
        if kind == "idpel_tidak_ketemu"
        else ",-6.2,106.8\n,-6.2,106.8\n"
    )
    target = upload(client, kind, "IDPEL,LAT,LON\n" + targets)
    result = client.post(
        f"/corrections/{kind}/checks",
        json={
            "batch_id": target,
            "pdpj_review_id": review["id"],
            "radius_meters": 50,
            "actor": "Uji",
            "request_key": "all_observations_test",
            "incremental": True,
        },
    )
    assert result.status_code == 200, result.json()
    return result.json()


@pytest.mark.parametrize("kind", ["idpel_kosong", "idpel_tidak_ketemu"])
def test_search_all_sources_then_deduplicate_and_export_selected_observation(
    client: TestClient,
    db_session: Session,
    kind: str,
) -> None:
    check = scenario(client, kind)
    # Progress helper uses the empty category; requests remain explicitly scoped here.
    base = f"/corrections/{kind}/checks/{check['process_id']}"
    run = check["latest_run"]
    for _ in range(20):
        if run["status"] == "completed":
            break
        response = client.post(
            base + "/advance",
            json={
                "run_id": run["id"],
                "revision": run["processing"]["revision"],
                "chunk_size": 1,
            },
        )
        assert response.status_code == 200, response.json()
        run = response.json()
    assert run["status"] == "completed"
    assert run["candidate_policy"] == 2
    rows = client.get(
        f"/corrections/{kind}/records", params={"process_id": check["process_id"]}
    ).json()["items"]
    detail = client.get(
        f"/corrections/{kind}/records/{rows[0]['id']}",
        params={"process_id": check["process_id"]},
    ).json()
    assert len(detail["candidates"]) == 2
    assert [item["fields"]["idpel"] for item in detail["candidates"]] == ["003", "002"]
    selected = detail["candidates"][1]
    assert selected["fields"]["lat"] == -6.2001
    assert selected["fields"]["coordinate_max_difference_meters"] > 1000
    assert len(selected["observations"]) == 3
    evidence = next(
        item
        for item in selected["observations"]
        if item["record_id"] == selected["fields"]["source_record_id"]
    )
    assert evidence["dataset_type"] == "dil"
    assert evidence["fields"]["lat"] == selected["fields"]["lat"]
    assert detail["candidates"][0]["fields"]["lat"] == -6.20005
    assert db_session.query(MatchCandidate).count() == 4
    assert db_session.query(CorrectionCandidateEvidence).count() == 4
    assert db_session.query(CorrectionMasterObservation).count() == 6
    choice = client.post(
        base + f"/records/{rows[0]['id']}/choice",
        json={
            "generation": 1,
            "expected_revision": 0,
            "candidate_id": selected["id"],
            "actor": "Uji",
            "reason": "Periksa sumber",
            "request_key": "observation_choice",
        },
    )
    assert choice.status_code == 200, choice.json()
    response = client.get(base + "/results/export")
    assert response.status_code == 200, (
        response.text[:100] if response.status_code != 200 else ""
    )
    book = load_workbook(BytesIO(response.content), read_only=True)
    exported = list(book.worksheets[0].values)
    assert exported[1][1] == "002"
    assert exported[1][-2:] == ("106.8", "-6.2001")
    assert "dil.csv" in exported[1][6]


def test_policy_migration_preserves_existing_runs_and_is_idempotent() -> None:
    engine = create_engine("sqlite:///:memory:")
    with engine.begin() as connection:
        connection.exec_driver_sql(
            "CREATE TABLE matching_runs (id INTEGER PRIMARY KEY)"
        )
        connection.exec_driver_sql("INSERT INTO matching_runs VALUES (1)")
        upgrade_correction_observations(connection)
        upgrade_correction_observations(connection)
        assert (
            connection.exec_driver_sql(
                "SELECT candidate_policy FROM matching_runs"
            ).scalar_one()
            == 1
        )
    engine.dispose()


def test_new_observation_search_preserves_previously_approved_coordinate(
    client: TestClient,
    db_session: Session,
) -> None:
    check = scenario(client, latest_lat="-6.2002")
    old = db_session.get(MatchingRun, check["latest_run"]["id"])
    assert old is not None
    old.candidate_policy = 1
    db_session.commit()
    completed = finish(client, check, check["latest_run"])
    rows = client.get(
        "/corrections/idpel_kosong/records", params={"process_id": check["process_id"]}
    ).json()["items"]
    record_id = rows[0]["id"]
    old_detail = client.get(
        f"/corrections/idpel_kosong/records/{record_id}",
        params={"process_id": check["process_id"]},
    ).json()
    chosen = old_detail["candidates"][0]
    assert chosen["fields"]["lat"] == -6.2002
    response = client.post(
        path(check) + f"/records/{record_id}/choice",
        json={
            "generation": 1,
            "expected_revision": 0,
            "candidate_id": chosen["id"],
            "actor": "Uji",
            "request_key": "preserved_observation_choice",
        },
    )
    assert response.status_code == 200, response.json()
    newer = client.post(
        path(check) + "/retry",
        json={
            "latest_run_id": completed["id"],
            "radius_meters": 50,
            "actor": "Uji",
            "request_key": "new_observation_version",
            "incremental": True,
        },
    ).json()
    finish(client, newer, newer["latest_run"])
    detail = client.get(
        f"/corrections/idpel_kosong/records/{record_id}",
        params={"process_id": check["process_id"]},
    ).json()
    same_id = next(
        item for item in detail["candidates"] if item["fields"]["idpel"] == "002"
    )
    assert same_id["fields"]["lat"] == -6.2001
    choice = client.get(path(check) + f"/records/{record_id}/choice").json()
    assert choice["current"]["candidate_id"] == chosen["id"]
    assert choice["current"]["candidate"]["fields"]["lat"] == -6.2002
    book = load_workbook(
        BytesIO(client.get(path(check) + "/results/export").content), read_only=True
    )
    assert list(book["Hasil Koreksi"].values)[1][-1] == "-6.2002"


def test_legacy_checkpoint_resumes_old_policy_and_explicit_retry_starts_from_zero(
    client: TestClient,
    db_session: Session,
) -> None:
    check = scenario(client)
    old = db_session.get(MatchingRun, check["latest_run"]["id"])
    assert old is not None
    old.candidate_policy = 1
    db_session.commit()
    run = advance(client, check, check["latest_run"])
    assert run["processing"]["stage"] == "matching"
    # Simulate an existing checkpoint saved by the previous release.
    response = client.post(
        path(check) + "/advance",
        json={
            "run_id": run["id"],
            "revision": run["processing"]["revision"],
            "chunk_size": 1,
        },
    )
    first = response.json()
    assert first["processing"]["processed_rows"] == 1
    assert first["processing"]["matched_rows"] == 0
    paused = client.post(path(check) + "/pause", json={"run_id": run["id"]}).json()
    assert paused["processing"]["state"] == "paused"
    resumed = client.post(path(check) + "/resume", json={"run_id": run["id"]}).json()
    assert resumed["candidate_policy"] == 1
    assert (
        resumed["processing"]["last_source_id"] == first["processing"]["last_source_id"]
    )
    body = {
        "latest_run_id": run["id"],
        "radius_meters": 50,
        "actor": "Uji",
        "request_key": "upgrade_legacy_check",
        "incremental": True,
    }
    assert client.post(path(check) + "/retry", json=body).status_code == 409
    paused = client.post(path(check) + "/pause", json={"run_id": run["id"]}).json()
    assert paused["processing"]["state"] == "paused"
    response = client.post(path(check) + "/retry", json=body)
    assert response.status_code == 200, response.json()
    new_check = response.json()
    new = new_check["latest_run"]
    assert new["candidate_policy"] == 2 and new["generation"] == 2
    assert (
        new["processing"]["processed_rows"] == new["processing"]["last_source_id"] == 0
    )
    assert (
        client.post(path(check) + "/retry", json=body).json()["latest_run"]["id"]
        == new["id"]
    )
    completed = finish(client, new_check, new)
    assert completed["matched_targets"] == 2
    old_outcomes = db_session.query(CorrectionTarget).filter_by(run_id=old.id).all()
    assert len(old_outcomes) == 1 and old_outcomes[0].search_status == "no_candidates"
    assert (
        client.post(path(check) + "/resume", json={"run_id": old.id}).status_code == 409
    )
    catalog_count = db_session.query(CorrectionMasterObservation).count()
    new_body = {
        **body,
        "latest_run_id": completed["id"],
        "request_key": "reuse_observation_index",
    }
    third = client.post(path(check) + "/retry", json=new_body).json()
    prepared = advance(client, third, third["latest_run"])
    assert prepared["processing"]["preparation"]["observation_rows"] == catalog_count
    assert db_session.query(CorrectionMasterObservation).count() == catalog_count


def test_observation_preparation_checkpoints_resume_without_duplicate_rows(
    client: TestClient,
    db_session: Session,
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    check = scenario(client)
    monkeypatch.setattr(observations, "OBSERVATION_CHUNK_ROWS", 2)
    run = advance(client, check, check["latest_run"])
    assert run["processing"]["preparation"]["phase"] == "observations"
    count = db_session.query(CorrectionMasterObservation).count()
    paused = client.post(path(check) + "/pause", json={"run_id": run["id"]}).json()
    restored = client.get(path(check) + "/progress").json()
    assert restored["processing"] == paused["processing"]
    assert db_session.query(CorrectionMasterObservation).count() == count
    run = client.post(path(check) + "/resume", json={"run_id": run["id"]}).json()
    run = finish(client, check, run)
    assert run["matched_targets"] == 2
    assert db_session.query(CorrectionObservationCatalog).one().processing["ready"]
    assert db_session.query(CorrectionMasterObservation).count() == 6
