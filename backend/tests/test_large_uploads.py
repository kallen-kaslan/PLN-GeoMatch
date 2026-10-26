"""Large uploads stay bounded, atomic, and preserve source data across chunks."""

import codecs
import io
import re
from contextlib import contextmanager
from tempfile import TemporaryFile
from zipfile import ZipFile

import pytest
from fastapi import HTTPException
from openpyxl import Workbook
from sqlalchemy.exc import SQLAlchemyError

from app.models.upload import SourceRecord, UploadBatch
from app.modules.upload import reader, router


def post_csv(client, content):
    return client.post(
        "/upload/",
        data={"dataset_type": "dil"},
        files={"file": ("large.csv", content, "text/csv")},
    )


def test_real_500_mb_file_size_boundary():
    # Sparse file checks the actual production boundary without allocating 500 MB RAM.
    with TemporaryFile() as source:
        source.truncate(500 * 1024 * 1024)
        reader.check_file_size(source)
        assert source.tell() == 0
        source.truncate(500 * 1024 * 1024 + 1)
        with pytest.raises(HTTPException) as error:
            reader.check_file_size(source)
        assert error.value.status_code == 413


def test_http_size_limit_accepts_exact_limit_and_rejects_one_byte_over(
    client, db_session, monkeypatch
):
    content = b"IDPEL,NAMA\n001,Nama\n"
    monkeypatch.setattr(reader, "MAX_UPLOAD_BYTES", len(content))
    assert post_csv(client, content).status_code == 200
    rejected = post_csv(client, content + b"\n")
    assert rejected.status_code == 413
    assert "500 MB" in rejected.json()["detail"]
    assert db_session.query(UploadBatch).count() == 1
    assert db_session.query(SourceRecord).count() == 1
    assert post_csv(client, content).status_code == 200


def test_chunked_upload_and_mapping_preserve_all_rows(client, db_session, monkeypatch):
    monkeypatch.setattr(reader, "IMPORT_CHUNK_ROWS", 3)
    monkeypatch.setattr(router, "IMPORT_CHUNK_ROWS", 3)
    sizes = []
    original_insert = db_session.bulk_insert_mappings

    def insert(model, mappings):
        sizes.append(len(mappings))
        return original_insert(model, mappings)

    monkeypatch.setattr(db_session, "bulk_insert_mappings", insert)
    content = "\ufeffIDPEL;NAMA;;NAMA\n" + "".join(
        f"00{index};Pelanggan {index};;NA\n" for index in range(8)
    )
    response = post_csv(client, content.encode("utf-8"))
    assert response.status_code == 200, response.text
    assert response.json()["total_rows"] == 8
    assert sizes == [3, 3, 2]
    batch_id = response.json()["id"]
    mapped = client.post(
        "/upload/confirm-mapping",
        json={"batch_id": batch_id, "column_mapping": {"idpel": "IDPEL"}},
    )
    assert mapped.status_code == 200, mapped.text
    assert mapped.json()["valid_rows"] == 8
    records = db_session.query(SourceRecord).order_by(SourceRecord.id).all()
    assert [record.row_number for record in records] == list(range(2, 10))
    assert [record.idpel for record in records] == [f"00{index}" for index in range(8)]
    assert records[-1].data == {
        "IDPEL": "007",
        "NAMA": "Pelanggan 7",
        "Kolom_3": "",
        "NAMA__2": "NA",
    }
    assert records[-1].normalized_data["idpel"] == "007"


def test_late_parse_failure_rolls_back_already_inserted_chunks(
    client, db_session, monkeypatch
):
    monkeypatch.setattr(reader, "IMPORT_CHUNK_ROWS", 2)
    monkeypatch.setattr(router, "IMPORT_CHUNK_ROWS", 2)
    response = post_csv(client, b'IDPEL,NAMA\n001,A\n002,B\n003,C\n004,"unterminated\n')
    assert response.status_code == 400
    assert db_session.query(UploadBatch).count() == 0
    assert db_session.query(SourceRecord).count() == 0
    assert post_csv(client, b"IDPEL\n005\n").status_code == 200


def test_late_database_failure_rolls_back_entire_upload(
    client, db_session, monkeypatch
):
    monkeypatch.setattr(router, "IMPORT_CHUNK_ROWS", 2)
    original_insert = db_session.bulk_insert_mappings
    calls = 0

    def insert(model, mappings):
        nonlocal calls
        calls += 1
        if calls == 2:
            raise SQLAlchemyError("insert failed")
        return original_insert(model, mappings)

    monkeypatch.setattr(db_session, "bulk_insert_mappings", insert)
    response = post_csv(client, b"IDPEL\n001\n002\n003\n")
    assert response.status_code == 500
    assert db_session.query(UploadBatch).count() == 0
    assert db_session.query(SourceRecord).count() == 0


@pytest.mark.parametrize("content", [b"", b"IDPEL\n", b"IDPEL\n\x81\n"])
def test_invalid_or_empty_upload_leaves_no_batch(client, db_session, content):
    assert post_csv(client, content).status_code == 400
    assert db_session.query(UploadBatch).count() == 0


@pytest.mark.parametrize(
    "encoding, prefix",
    [
        ("utf-8", b""),
        ("utf-8-sig", b""),
        ("cp1252", b""),
        ("utf-16", b""),
        ("utf-16-be", codecs.BOM_UTF16_BE),
    ],
)
def test_csv_encoding_preserves_special_characters_and_identifiers(
    client, db_session, encoding, prefix
):
    text = "IDPEL;NAMA;ALAMAT\n000123;Ìsmi – Café;Jalan François\n"
    uploaded = post_csv(client, prefix + text.encode(encoding))
    assert uploaded.status_code == 200, uploaded.text
    record = db_session.query(SourceRecord).one()
    assert record.data == {
        "IDPEL": "000123",
        "NAMA": "Ìsmi – Café",
        "ALAMAT": "Jalan François",
    }


def test_csv_detects_windows_encoding_beyond_delimiter_sample(client, db_session):
    content = (
        b"IDPEL,NAMA\n"
        + b"001,Nama panjang dari ekspor Excel\n" * 1200
        + b"002,\xccsmi\n"
    )
    assert content.index(b"\xcc") > 32768
    uploaded = post_csv(client, content)
    assert uploaded.status_code == 200, uploaded.text
    assert uploaded.json()["total_rows"] == 1201
    last = db_session.query(SourceRecord).order_by(SourceRecord.id.desc()).first()
    assert last.data["NAMA"] == "Ìsmi"


def test_csv_with_corrupt_bom_encoding_does_not_silently_replace_characters(
    client, db_session
):
    uploaded = post_csv(client, codecs.BOM_UTF8 + b"IDPEL,NAMA\n001,\xcc\n")
    assert uploaded.status_code == 400
    assert db_session.query(UploadBatch).count() == 0


def test_excel_reads_selected_sheet_with_internal_blank_rows(client, db_session):
    book = Workbook()
    book.active.title = "Instructions"
    sheet = book.create_sheet("Data")
    sheet.append(["IDPEL", "NAMA"])
    sheet.append(["001", "Satu"])
    sheet.append([None, None])
    sheet.append(["002", "Dua"])
    sheet.append([None, None])
    content = io.BytesIO()
    book.save(content)
    response = client.post(
        "/upload/",
        data={"dataset_type": "dil", "sheet_name": "Data"},
        files={"file": ("large.xlsx", content.getvalue())},
    )
    assert response.status_code == 200, response.text
    assert response.json()["total_rows"] == 3
    records = db_session.query(SourceRecord).order_by(SourceRecord.id).all()
    assert records[1].data == {"IDPEL": "", "NAMA": ""}
    assert records[-1].row_number == 4
    assert db_session.query(UploadBatch).one().source_sheet == "Data"


@pytest.mark.parametrize(
    "header_row, leading_value", [(2, None), (2, "   "), (8, "   "), (70, "   ")]
)
def test_excel_skips_leading_empty_rows_and_preserves_source_numbers(
    client, db_session, header_row, leading_value
):
    book = Workbook()
    sheet = book.active
    sheet.title = "DIL_SALDO_MASK_UP_202606_53321"
    sheet.cell(1, 1, leading_value)
    for column, value in enumerate(["IDPEL", "NAMA", "", "NAMA"], start=1):
        sheet.cell(header_row, column, value)
    for column, value in enumerate(["000123", "Satu", "", "Asli"], start=1):
        sheet.cell(header_row + 1, column, value)
    for column, value in enumerate(["000456", "Dua", "", "Asli dua"], start=1):
        sheet.cell(header_row + 3, column, value)
    sheet.cell(header_row + 4, 1, "")
    content = io.BytesIO()
    book.save(content)
    uploaded = client.post(
        "/upload/",
        data={"dataset_type": "dil", "sheet_name": sheet.title},
        files={"file": ("dil.xlsx", content.getvalue())},
    )
    assert uploaded.status_code == 200, uploaded.text
    assert uploaded.json()["total_rows"] == 3
    batch_id = uploaded.json()["id"]
    preview = client.get(f"/master/dil/batches/{batch_id}").json()
    assert preview["headers"] == ["IDPEL", "NAMA", "Kolom_3", "NAMA__2"]
    assert preview["row_numbers"] == list(range(header_row + 1, header_row + 4))
    records = db_session.query(SourceRecord).order_by(SourceRecord.id).all()
    assert [record.row_number for record in records] == list(
        range(header_row + 1, header_row + 4)
    )
    assert records[0].data["IDPEL"] == "000123"
    assert records[1].data == {"IDPEL": "", "NAMA": "", "Kolom_3": "", "NAMA__2": ""}
    checked = client.post(
        "/upload/confirm-mapping",
        json={"batch_id": batch_id, "column_mapping": {"idpel": "IDPEL"}},
    )
    assert checked.status_code == 200, checked.text
    assert checked.json()["valid_rows"] == 2
    assert checked.json()["errors"][0]["row"] == header_row + 2


def test_excel_without_nonempty_header_is_rejected_atomically(client, db_session):
    book = Workbook()
    book.active.cell(1, 1, "   ")
    book.active.cell(8, 2, "")
    content = io.BytesIO()
    book.save(content)
    uploaded = client.post(
        "/upload/",
        data={"dataset_type": "dil"},
        files={"file": ("empty.xlsx", content.getvalue())},
    )
    assert uploaded.status_code == 400
    assert db_session.query(UploadBatch).count() == 0


@pytest.mark.parametrize("header_row", [1, 2, 70])
def test_excel_incorrect_dimensions_preview_and_import_actual_cells(
    client, db_session, header_row
):
    book = Workbook()
    sheet = book.active
    sheet.title = "DIL_SALDO_MASK_UP_202606_53311"
    for column, value in enumerate(["IDPEL", "NAMA", "", "NAMA"], start=1):
        sheet.cell(header_row, column, value)
    for offset in range(8):
        sheet.cell(header_row + offset + 1, 1, f"00012{offset}")
        sheet.cell(header_row + offset + 1, 2, f"Contoh {offset}")
    original = io.BytesIO()
    book.save(original)
    patched = io.BytesIO()
    with (
        ZipFile(io.BytesIO(original.getvalue())) as source,
        ZipFile(patched, "w") as target,
    ):
        for item in source.infolist():
            data = source.read(item.filename)
            if item.filename == "xl/worksheets/sheet1.xml":
                data = re.sub(
                    rb'<dimension ref="[^"]+"', b'<dimension ref="A1:A1"', data
                )
            target.writestr(item, data)
    file = {"file": ("dil.xlsx", patched.getvalue())}
    preview = client.post(
        "/upload/preview", data={"sheet_name": sheet.title}, files=file
    )
    assert preview.status_code == 200, preview.text
    expected_headers = ["IDPEL", "NAMA", "Kolom_3", "NAMA__2"]
    assert preview.json()["headers"] == expected_headers
    assert preview.json()["headerRow"] == header_row
    assert len(preview.json()["rows"]) == 5
    assert preview.json()["rows"][0] == {
        "IDPEL": "000120",
        "NAMA": "Contoh 0",
        "Kolom_3": "",
        "NAMA__2": "",
    }
    assert db_session.query(UploadBatch).count() == 0
    assert db_session.query(SourceRecord).count() == 0
    uploaded = client.post(
        "/upload/", data={"dataset_type": "dil", "sheet_name": sheet.title}, files=file
    )
    assert uploaded.status_code == 200, uploaded.text
    stored = client.get(f"/master/dil/batches/{uploaded.json()['id']}").json()
    assert uploaded.json()["total_rows"] == 8
    assert stored["headers"] == expected_headers
    assert stored["preview"] == preview.json()["rows"]
    assert stored["row_numbers"] == list(range(header_row + 1, header_row + 6))


def test_preview_enforces_size_and_reports_sheet_error_without_saving(
    client, db_session, monkeypatch
):
    book = Workbook()
    book.active.title = "Kosong"
    content = io.BytesIO()
    book.save(content)
    files = {"file": ("empty.xlsx", content.getvalue())}
    empty = client.post("/upload/preview", data={"sheet_name": "Kosong"}, files=files)
    assert empty.status_code == 400 and "Kosong" in empty.json()["detail"]
    missing = client.post("/upload/preview", data={"sheet_name": "DIL"}, files=files)
    assert missing.status_code == 400 and "tidak ditemukan" in missing.json()["detail"]
    monkeypatch.setattr(reader, "MAX_UPLOAD_BYTES", len(content.getvalue()) - 1)
    assert client.post("/upload/preview", files=files).status_code == 413
    assert db_session.query(UploadBatch).count() == 0
    assert db_session.query(SourceRecord).count() == 0


def test_upload_does_not_read_entire_file_into_bytes(client, monkeypatch):
    original_read_upload = router.read_upload

    @contextmanager
    def checked_read_upload(source, filename, sheet):
        class BoundedRead:
            def read(self, size=-1):
                assert size >= 0, "Unbounded reads would load the whole upload into RAM"
                return source.read(size)

            def __getattr__(self, name):
                return getattr(source, name)

        with original_read_upload(BoundedRead(), filename, sheet) as result:
            yield result

    monkeypatch.setattr(router, "read_upload", checked_read_upload)
    assert post_csv(client, b"IDPEL\n001\n").status_code == 200
