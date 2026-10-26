"""Original workbook history, explicit sheet confirmation, and reference isolation."""

from io import BytesIO
from pathlib import Path

import pytest
from fastapi.testclient import TestClient
from openpyxl import Workbook
from sqlalchemy import inspect
from sqlalchemy.orm import Session

from app.config import settings
from app.models.upload import SourceRecord, UploadBatch
from app.models.upload_file import UploadFileArchive, UploadSheetImport
from app.modules.upload.archive import archive_path
from app.upload_file_schema import upgrade_upload_files


def workbook() -> bytes:
    book = Workbook()
    book.active.title = "Januari"
    book.active.append(["IDPEL", "NAMA", "LAT", "LON"])
    book.active.append(["001", "Pelanggan A", -6.2, 106.8])
    sheet = book.create_sheet("Februari")
    sheet.append(["IDPEL", "NAMA", "LAT", "LON"])
    sheet.append(["002", "Pelanggan B", -6.3, 106.9])
    book.create_sheet("Kosong")
    output = BytesIO()
    book.save(output)
    return output.getvalue()


def archive(
    client: TestClient, content: bytes | None = None, filename: str = "multi.xlsx"
) -> dict:
    response = client.post(
        "/upload/files",
        files={"file": (filename, content if content is not None else workbook())},
    )
    assert response.status_code == 200, response.text
    return response.json()


def import_sheet(
    client: TestClient, file_id: int, sheet: str | None, **extra: str
) -> dict:
    response = client.post(
        f"/upload/files/{file_id}/import-sheet",
        json={"dataset_type": "dil", "sheet_name": sheet, **extra},
    )
    assert response.status_code == 200, response.text
    return response.json()


def confirm(client: TestClient, batch: int) -> None:
    response = client.post(
        "/upload/confirm-mapping",
        json={"batch_id": batch, "column_mapping": {"idpel": "IDPEL"}},
    )
    assert response.status_code == 200, response.text


def test_original_and_previews_do_not_create_reference_or_batches(
    client: TestClient, db_session: Session
) -> None:
    stored = archive(client)
    assert stored["sheet_names"] == ["Januari", "Februari", "Kosong"]
    assert stored["imports"] == []
    for name, expected in [("Januari", "001"), ("Februari", "002")]:
        preview = client.get(
            f"/upload/files/{stored['id']}/preview", params={"sheet_name": name}
        )
        assert preview.status_code == 200, preview.text
        assert preview.json()["rows"][0]["IDPEL"] == expected
    assert (
        db_session.query(UploadBatch).count()
        == db_session.query(SourceRecord).count()
        == 0
    )
    assert client.get("/master/dil/summary").json()["total_rows"] == 0
    assert client.get("/upload/files").json() == [stored]


def test_sheet_status_changes_only_after_mapping_confirmation_and_no_duplicates(
    client: TestClient, db_session: Session
) -> None:
    stored = archive(client)
    first = import_sheet(client, stored["id"], "Januari")
    assert first["status"] == "needs_mapping"
    assert client.get("/master/dil/summary").json()["total_rows"] == 0
    assert (
        client.post(
            "/upload/confirm-mapping",
            json={"batch_id": first["id"], "column_mapping": {"idpel": "SALAH"}},
        ).status_code
        == 422
    )
    assert (
        client.get("/upload/files").json()[0]["imports"][0]["status"] == "needs_mapping"
    )
    assert import_sheet(client, stored["id"], "Januari")["id"] == first["id"]
    confirm(client, first["id"])
    assert client.get("/master/dil/summary").json()["total_rows"] == 1
    history = client.get("/upload/files").json()[0]
    assert (
        len(history["imports"]) == 1 and history["imports"][0]["status"] == "completed"
    )
    # Reopen the same original file and import a different sheet without another multipart upload.
    second = import_sheet(client, stored["id"], "Februari")
    assert second["id"] != first["id"]
    assert client.get("/master/dil/summary").json()["total_rows"] == 1
    confirm(client, second["id"])
    assert client.get("/master/dil/summary").json()["total_rows"] == 2
    assert import_sheet(client, stored["id"], "Januari")["status"] == "completed"
    assert db_session.query(UploadBatch).count() == 2
    assert db_session.query(SourceRecord).count() == 2
    assert (
        client.post(
            "/upload/confirm-mapping",
            json={"batch_id": first["id"], "column_mapping": {"idpel": "NAMA"}},
        ).status_code
        == 409
    )


def test_fingerprint_deduplicates_reuploads_but_not_changed_contents(
    client: TestClient, db_session: Session
) -> None:
    content = workbook()
    first = archive(client, content)
    assert archive(client, content, "renamed.xlsx")["id"] == first["id"]
    changed = archive(client, b"IDPEL\n003\n", "multi.csv")
    assert changed["id"] != first["id"]
    assert len(list(Path(settings.UPLOAD_STORAGE_DIR).iterdir())) == 2
    assert db_session.query(UploadFileArchive).count() == 2


def test_scope_per_kind_period_csv_and_draft_reuse(client: TestClient) -> None:
    file = archive(client, b"IDPEL,NAMA\n001,A\n", "data.csv")
    assert file["sheet_names"] == []
    first = import_sheet(client, file["id"], None, period="2026-01")
    confirm(client, first["id"])
    assert import_sheet(client, file["id"], None, period="2026-01")["id"] == first["id"]
    another_period = import_sheet(client, file["id"], None, period="2026-02")
    assert another_period["id"] != first["id"]
    another_kind = client.post(
        f"/upload/files/{file['id']}/import-sheet",
        json={"dataset_type": "pdpj", "period": "2026-01"},
    )
    assert (
        another_kind.status_code == 200
        and another_kind.json()["status"] == "needs_mapping"
    )
    assert len(client.get("/upload/files").json()[0]["imports"]) == 3
    assert (
        client.get(f"/upload/files/{file['id']}/preview").json()["rows"][0]["IDPEL"]
        == "001"
    )
    assert (
        client.get(f"/upload/files/{file['id']}/preview?sheet_name=Januari").status_code
        == 422
    )


def test_invalid_sheet_and_missing_original_leave_no_partial_reference(
    client: TestClient, db_session: Session
) -> None:
    content = workbook()
    stored = archive(client, content)
    url = f"/upload/files/{stored['id']}"
    for name in [None, "../Januari", "Hilang"]:
        assert (
            client.post(
                url + "/import-sheet", json={"dataset_type": "dil", "sheet_name": name}
            ).status_code
            == 422
        )
    assert client.get(url + "/preview?sheet_name=Kosong").status_code == 400
    assert (
        client.post(
            url + "/import-sheet", json={"dataset_type": "dil", "sheet_name": "Kosong"}
        ).status_code
        == 400
    )
    assert (
        db_session.query(UploadBatch).count()
        == db_session.query(UploadSheetImport).count()
        == 0
    )
    entry = db_session.get(UploadFileArchive, stored["id"])
    path = archive_path(entry.storage_key)
    path.unlink()
    assert client.get(url + "/preview?sheet_name=Januari").status_code == 410
    assert (
        client.post(
            url + "/import-sheet", json={"dataset_type": "dil", "sheet_name": "Januari"}
        ).status_code
        == 410
    )
    # Identical reupload repairs the original, retaining its history identity.
    recovered = archive(client, content)
    assert recovered["id"] == stored["id"]
    assert (
        client.get(
            f"/upload/files/{recovered['id']}/preview?sheet_name=Januari"
        ).status_code
        == 200
    )
    assert client.get("/upload/files/99999/preview").status_code == 404


def test_reset_and_migration_preserve_foreign_key_integrity(
    client: TestClient, db_session: Session
) -> None:
    stored = archive(client)
    imported = import_sheet(client, stored["id"], "Januari")
    confirm(client, imported["id"])
    paths = list(Path(settings.UPLOAD_STORAGE_DIR).iterdir())
    assert client.post("/upload/reset").status_code == 200
    assert client.get("/upload/files").json() == []
    assert all(not path.exists() for path in paths)
    connection = db_session.connection()
    UploadSheetImport.__table__.drop(connection)
    UploadFileArchive.__table__.drop(connection)
    upgrade_upload_files(connection)
    upgrade_upload_files(connection)
    assert "upload_file_archives" in inspect(connection).get_table_names()


def test_invalid_uploads_are_not_archived(
    client: TestClient, db_session: Session
) -> None:
    for filename, content in [
        ("invalid.xlsx", b"invalid"),
        ("file.txt", b"IDPEL\n001"),
    ]:
        assert (
            client.post(
                "/upload/files", files={"file": (filename, content)}
            ).status_code
            == 400
        )
    assert db_session.query(UploadFileArchive).count() == 0
    with pytest.raises(ValueError):
        archive_path("../invalid")


def test_failed_sheet_write_rolls_back_rows_but_retains_original_for_retry(
    client: TestClient, db_session: Session, monkeypatch: pytest.MonkeyPatch
) -> None:
    from sqlalchemy.exc import SQLAlchemyError

    from app.modules.upload import ingestion

    stored = archive(client, b"IDPEL\n001\n002\n003\n", "retry.csv")
    monkeypatch.setattr(ingestion, "IMPORT_CHUNK_ROWS", 2)
    original = db_session.bulk_insert_mappings
    calls = 0

    def insert(model, values):
        nonlocal calls
        calls += 1
        if calls == 2:
            raise SQLAlchemyError("simulated write failure")
        return original(model, values)

    monkeypatch.setattr(db_session, "bulk_insert_mappings", insert)
    response = client.post(
        f"/upload/files/{stored['id']}/import-sheet", json={"dataset_type": "dil"}
    )
    assert response.status_code == 500
    assert (
        db_session.query(UploadBatch).count()
        == db_session.query(SourceRecord).count()
        == 0
    )
    assert client.get("/upload/files").json()[0]["imports"] == []
    monkeypatch.setattr(db_session, "bulk_insert_mappings", original)
    batch = import_sheet(client, stored["id"], None)
    assert batch["status"] == "needs_mapping"
    assert db_session.query(SourceRecord).count() == 3
