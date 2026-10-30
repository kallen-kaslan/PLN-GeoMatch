"""Saved choices project immediately into partial results, history and export."""

import io

import pytest
from fastapi.testclient import TestClient
from openpyxl import Workbook, load_workbook
from sqlalchemy.orm import Session

from app.models.correction_decision import CorrectionDecision
from app.models.upload import SourceRecord, UploadBatch
from app.modules.corrections.spreadsheet import RESULT_HEADERS, XLSX_MEDIA_TYPE
from tests.test_correction_map import body
from tests.test_corrections import upload


def sheet_rows(book, name: str) -> list[dict]:
    rows = iter(book[name].values)
    headers = next(rows)
    return [dict(zip(headers, row)) for row in rows]


@pytest.fixture(params=["idpel_kosong", "idpel_tidak_ketemu"])
def partial(client: TestClient, request: pytest.FixtureRequest) -> dict:
    kind = request.param
    master = Workbook()
    sheet = master.active
    sheet.title = "Master Pelanggan"
    sheet.append(["IDPEL", "NAMA", "LAT", "LON", "JENISPELANGGAN"])
    sheet.append(["001", "PDPJ", -6.2, 106.8, "PRABAYAR"])
    sheet.append(["002", "First", -6.200005, 106.8, "PRABAYAR"])
    sheet.append(["003", "Second", -6.20001, 106.8, "PRABAYAR"])
    content = io.BytesIO()
    master.save(content)
    imported = client.post(
        "/upload/",
        data={"dataset_type": "dil", "sheet_name": sheet.title},
        files={"file": ("master-pelanggan.xlsx", content.getvalue(), XLSX_MEDIA_TYPE)},
    )
    assert imported.status_code == 200
    assert (
        client.post(
            "/upload/confirm-mapping",
            json={
                "batch_id": imported.json()["id"],
                "column_mapping": {"idpel": "IDPEL"},
            },
        ).status_code
        == 200
    )
    pdpj = upload(client, "pdpj", "IDPEL,LAT,LON\n001,-6.2,106.8\n")
    review = client.post("/pdpj/reviews", json={"batch_id": pdpj}).json()
    assert (
        client.post(
            f"/pdpj/reviews/{review['id']}/complete",
            json={"revision": review["revision"], "actor": "Petugas"},
        ).status_code
        == 200
    )
    batch = upload(
        client,
        kind,
        "IDPEL,NAMA,LAT,LON,NOAGENDA,VERIF,JENISPELANGGAN\n"
        + "".join(
            f"{700 + i if kind == 'idpel_tidak_ketemu' else ''},Target {i},-6.2,106.8,000AGENDA{i},202606,PRABAYAR\n"
            for i in range(5)
        ),
    )
    checked = client.post(
        f"/corrections/{kind}/checks",
        json={
            "batch_id": batch,
            "pdpj_review_id": review["id"],
            "radius_meters": 10,
            "actor": "Petugas",
            "request_key": "result_request_1",
            "incremental": True,
        },
    )
    assert checked.status_code == 200, checked.json()
    result = checked.json()
    result["kind"] = kind
    result["base"] = f"/corrections/{kind}/checks/{result['process_id']}"
    run = result["latest_run"]
    for _ in range(2):
        run = advance(client, result, run)
    result["latest_run"] = run
    rows = client.get(
        f"/corrections/{kind}/records", params={"process_id": result["process_id"]}
    ).json()["items"]
    result["record"] = rows[0]["id"]
    result["candidates"] = client.get(
        f"/corrections/{kind}/records/{result['record']}",
        params={"process_id": result["process_id"]},
    ).json()["candidates"]
    return result


def advance(client: TestClient, check: dict, run: dict) -> dict:
    response = client.post(
        check["base"] + "/advance",
        json={
            "run_id": run["id"],
            "revision": run["processing"]["revision"],
            "chunk_size": 2,
        },
    )
    assert response.status_code == 200, response.json()
    return response.json()


def choose(
    client: TestClient,
    check: dict,
    index: int | None = 0,
    revision: int = 0,
    generation: int = 1,
) -> dict:
    candidate = check["candidates"][index]["id"] if index is not None else None
    response = client.post(
        check["base"] + f"/records/{check['record']}/choice",
        json=body(
            candidate,
            revision=revision,
            generation=generation,
            key=f"result_choice_{check['process_id']}_{revision}",
            reason="Bukti diperiksa" if revision else "",
        ),
    )
    assert response.status_code == 200, response.json()
    return response.json()


def test_partial_choice_appears_in_results_history_and_xlsx_without_changing_source(
    client: TestClient, db_session: Session, partial: dict
) -> None:
    source_before = dict(db_session.get(SourceRecord, partial["record"]).data)
    selected = choose(client, partial)
    result = client.get(partial["base"] + "/results")
    assert result.status_code == 200, result.json()
    value = result.json()
    assert value["total"] == 1
    summary = value["summary"]
    assert summary["total_targets"] == 5
    assert summary["saved_rows"] == 2
    assert summary["not_processed"] == 3
    assert summary["chosen_count"] == 1
    assert summary["not_selected"] == 4
    assert summary["partial"] is True
    assert value["items"][0]["choice"] == selected["current"]
    assert value["items"][0]["choice"]["candidate"]["fields"]["idpel"] == "002"
    assert client.get(partial["base"]).json()["current_generation"] == 0
    legacy_summary = client.get(f"/process/{partial['process_id']}/summary").json()
    assert legacy_summary["chosen_count"] == legacy_summary["valid_count"] == 1
    history = next(
        item
        for item in client.get("/history/processes").json()
        if item["id"] == partial["process_id"]
    )
    assert history["saved_rows"] == 2 and history["total_targets"] == 5
    assert history["chosen_count"] == 1
    events = client.get(partial["base"] + "/results/history").json()
    assert events["total"] == 1
    assert events["items"][0]["revision"] == 1
    assert events["items"][0]["row_number"] == 2
    for url in [
        partial["base"] + "/results/export",
        f"/process/{partial['process_id']}/export",
    ]:
        response = client.get(url)
        assert response.status_code == 200, response.text
        assert response.headers["content-type"] == XLSX_MEDIA_TYPE
        assert response.headers["content-disposition"].endswith('.xlsx"')
        book = load_workbook(io.BytesIO(response.content))
        sheet = book["Hasil Koreksi"]
        assert list(next(sheet.values)) == RESULT_HEADERS
        exported = sheet_rows(book, "Hasil Koreksi")
        assert len(exported) == 1
        row = exported[0]
        assert row["idpelanggan"] == "002"
        assert row["noagenda"] == "000AGENDA0"
        assert row["verif"] == "202606"
        assert row["jenispelanggan"] == "PRABAYAR"
        assert row["latitudey"] == "-6.2"
        assert row["longitudex"] == "106.8"
        assert row["sumber file (sheet)"] == "master-pelanggan.xlsx (Master Pelanggan)"
        assert row["sel (baris ke berapa)"] == "3"
        assert row["x_acc"] == "106.8"
        assert row["y_acc"] == "-6.200005"
        assert sheet_rows(book, "Keputusan Aktif")[0]["petugas"] == "Petugas Uji"
        assert sheet_rows(book, "Data Target")[0]["sel (baris ke berapa)"] == 2
        assert sheet_rows(book, "Observasi Sumber")[0]["koordinat_dipakai"] == "Ya"
        assert sheet.freeze_panes is None
        assert sheet.auto_filter.ref is None
        assert not sheet["A1"].font.bold
        assert sheet["A1"].fill.patternType is None
        assert sheet["A2"].fill.patternType is None
        assert sheet["I1"].value == "x_acc"
        assert sheet["J1"].value == "y_acc"
        assert sheet.column_dimensions["G"].width >= 50
        assert sheet["B2"].data_type == "s" and sheet["B2"].number_format == "@"
        assert sheet["E2"].data_type == "s"
        book.close()
    assert db_session.get(SourceRecord, partial["record"]).data == source_before
    wrong = (
        "idpel_kosong" if partial["kind"] != "idpel_kosong" else "idpel_tidak_ketemu"
    )
    assert (
        client.get(
            partial["base"].replace(partial["kind"], wrong) + "/results"
        ).status_code
        == 404
    )


def test_replacement_cancellation_filters_and_history_are_not_counted_as_extra_choices(
    client: TestClient, partial: dict
) -> None:
    choose(client, partial)
    choose(client, partial, index=1, revision=1)
    result = client.get(partial["base"] + "/results").json()
    assert result["total"] == result["summary"]["chosen_count"] == 1
    assert result["summary"]["decision_events"] == 2
    assert result["items"][0]["choice"]["candidate"]["fields"]["idpel"] == "003"
    assert (
        client.get(partial["base"] + "/results", params={"q": "Second"}).json()["total"]
        == 1
    )
    assert (
        client.get(partial["base"] + "/results", params={"q": "No match"}).json()[
            "total"
        ]
        == 0
    )
    assert (
        client.get(
            partial["base"] + "/results", params={"status": "not_processed"}
        ).json()["total"]
        == 3
    )
    choose(client, partial, index=None, revision=2)
    result = client.get(partial["base"] + "/results").json()
    assert result["total"] == result["summary"]["chosen_count"] == 0
    assert result["summary"]["cancelled_count"] == 1
    assert result["summary"]["decision_events"] == 3
    assert (
        client.get(partial["base"] + "/results", params={"status": "cancelled"}).json()[
            "total"
        ]
        == 1
    )
    assert (
        client.get(
            partial["base"] + "/results",
            params={"status": "all", "skip": 4, "limit": 1},
        ).json()["items"][0]["row_number"]
        == 6
    )
    events = client.get(partial["base"] + "/results/history").json()
    assert [item["revision"] for item in events["items"]] == [3, 2, 1]
    assert events["items"][0]["candidate_id"] is None
    book = load_workbook(
        io.BytesIO(client.get(partial["base"] + "/results/export").content)
    )
    assert sheet_rows(book, "Hasil Koreksi") == []
    book.close()
    history_file = client.get(partial["base"] + "/results/history/export")
    assert history_file.status_code == 200
    book = load_workbook(io.BytesIO(history_file.content))
    history = sheet_rows(book, "Riwayat Pilihan")
    assert [event["revisi_pilihan"] for event in history] == [3, 2, 1]
    assert [event["idpelanggan"] for event in history] == [None, "003", "002"]
    assert history[0]["status"] == "Pilihan dibatalkan"
    assert history[0]["x_acc"] is None
    assert history[0]["y_acc"] is None
    assert history[0]["sumber file (sheet)"] is None
    book.close()


def test_resume_and_new_search_preserve_choices_and_allow_historical_result_export(
    client: TestClient, partial: dict
) -> None:
    selected = choose(client, partial)
    run = partial["latest_run"]
    for _ in range(12):
        if run["status"] == "completed":
            break
        run = advance(client, partial, run)
    assert run["status"] == "completed"
    before = client.get(partial["base"] + "/results").json()
    assert before["summary"]["partial"] is False
    assert before["summary"]["saved_rows"] == 5
    assert before["items"][0]["choice"] == selected["current"]
    retried = client.post(
        partial["base"] + "/retry",
        json={
            "latest_run_id": run["id"],
            "radius_meters": 100,
            "actor": "Petugas",
            "request_key": "result_retry_1",
            "incremental": True,
        },
    )
    assert retried.status_code == 200, retried.json()
    run = retried.json()["latest_run"]
    preparing = client.get(partial["base"] + "/results").json()
    assert preparing["summary"]["generation"] == 2
    assert preparing["summary"]["saved_rows"] == 0
    assert preparing["items"][0]["choice"] == selected["current"]
    for _ in range(2):
        run = advance(client, partial, run)
    partial["candidates"] = client.get(
        f"/corrections/{partial['kind']}/records/{partial['record']}",
        params={"process_id": partial["process_id"]},
    ).json()["candidates"]
    choose(client, partial, index=1, revision=1, generation=2)
    current = client.get(partial["base"] + "/results").json()
    historical = client.get(
        partial["base"] + "/results", params={"generation": 1}
    ).json()
    assert current["items"][0]["choice"]["candidate"]["fields"]["idpel"] == "003"
    assert historical["items"][0]["choice"]["candidate"]["fields"]["idpel"] == "002"
    export = client.get(partial["base"] + "/results/export", params={"generation": 1})
    book = load_workbook(io.BytesIO(export.content))
    assert sheet_rows(book, "Hasil Koreksi")[0]["idpelanggan"] == "002"
    assert "radius_10m_v1.xlsx" in export.headers["content-disposition"]
    info = dict(list(book["Informasi"].values)[1:])
    assert info["Radius pemeriksaan (meter)"] == 10
    assert info["Versi pemeriksaan"] == 1
    assert sheet_rows(book, "Pemeriksaan Radius")[0]["radius_meter"] == 10
    assert sheet_rows(book, "Keputusan Aktif")[0]["radius_pilihan_meter"] == 10
    book.close()
    newer_file = client.get(
        partial["base"] + "/results/export", params={"generation": 2}
    )
    assert "radius_100m_v2.xlsx" in newer_file.headers["content-disposition"]
    newer_book = load_workbook(io.BytesIO(newer_file.content))
    assert sheet_rows(newer_book, "Hasil Koreksi")[0]["idpelanggan"] == "003"
    assert (
        dict(list(newer_book["Informasi"].values)[1:])["Radius pemeriksaan (meter)"]
        == 100
    )
    assert sheet_rows(newer_book, "Keputusan Aktif")[0]["radius_pilihan_meter"] == 100
    newer_book.close()
    historical_choice = client.get(
        partial["base"] + f"/records/{partial['record']}/choice",
        params={"generation": 1},
    ).json()
    assert historical_choice["current"]["candidate"]["fields"]["idpel"] == "002"
    historical_events = client.get(
        partial["base"] + "/results/history", params={"generation": 1}
    ).json()
    assert historical_events["total"] == 1
    history_file = client.get(
        partial["base"] + "/results/history/export", params={"generation": 1}
    )
    assert "radius_10m_v1.xlsx" in history_file.headers["content-disposition"]
    history_book = load_workbook(io.BytesIO(history_file.content))
    events = sheet_rows(history_book, "Riwayat Pilihan")
    assert len(events) == 1 and events[0]["idpelanggan"] == "002"
    assert events[0]["radius_pilihan_meter"] == 10
    history_book.close()
    assert (
        client.get(partial["base"] + "/results", params={"generation": 99}).status_code
        == 404
    )


def test_latest_choice_is_scoped_to_its_process_even_for_the_same_source_row(
    client: TestClient, partial: dict
) -> None:
    choose(client, partial)
    checked = client.post(
        f"/corrections/{partial['kind']}/checks",
        json={
            "batch_id": partial["batch_id"],
            "pdpj_review_id": partial["pdpj_review_id"],
            "radius_meters": 10,
            "actor": "Petugas",
            "request_key": "second_result_scope",
            "incremental": True,
        },
    ).json()
    other = {
        **checked,
        "base": f"/corrections/{partial['kind']}/checks/{checked['process_id']}",
        "record": partial["record"],
    }
    run = checked["latest_run"]
    for _ in range(2):
        run = advance(client, other, run)
    other["candidates"] = client.get(
        f"/corrections/{partial['kind']}/records/{partial['record']}",
        params={"process_id": checked["process_id"]},
    ).json()["candidates"]
    choose(client, other, index=1)
    first = client.get(partial["base"] + "/results").json()
    second = client.get(other["base"] + "/results").json()
    assert first["items"][0]["choice"]["candidate"]["fields"]["idpel"] == "002"
    assert second["items"][0]["choice"]["candidate"]["fields"]["idpel"] == "003"
    assert first["summary"]["chosen_count"] == second["summary"]["chosen_count"] == 1


def test_xlsx_keeps_formula_like_text_literal_without_changing_saved_decision(
    client: TestClient, partial: dict
) -> None:
    request = body(partial["candidates"][0]["id"], key="xlsx_formula_actor")
    request["actor"] = "=SUM(1,1)"
    url = partial["base"] + f"/records/{partial['record']}/choice"
    assert client.post(url, json=request).status_code == 200
    response = client.get(partial["base"] + "/results/export")
    book = load_workbook(io.BytesIO(response.content))
    row = sheet_rows(book, "Keputusan Aktif")[0]
    assert row["petugas"] == "=SUM(1,1)"
    assert book["Keputusan Aktif"]["E2"].data_type == "s"
    assert sheet_rows(book, "Hasil Koreksi")[0]["idpelanggan"] == "002"
    book.close()
    assert client.get(url).json()["current"]["actor"] == "=SUM(1,1)"


def test_xlsx_obeys_filters_and_history_record_scope(
    client: TestClient, partial: dict
) -> None:
    choose(client, partial)
    filtered = client.get(partial["base"] + "/results/export", params={"q": "No match"})
    book = load_workbook(io.BytesIO(filtered.content))
    assert sheet_rows(book, "Hasil Koreksi") == []
    book.close()
    pending = client.get(
        partial["base"] + "/results/export", params={"status": "not_processed"}
    )
    book = load_workbook(io.BytesIO(pending.content))
    rows = sheet_rows(book, "Hasil Koreksi")
    assert len(rows) == 3
    assert all(row["idpelanggan"] is None for row in rows)
    assert all(row["x_acc"] is None and row["y_acc"] is None for row in rows)
    book.close()
    history = client.get(
        partial["base"] + "/results/history/export",
        params={"record_id": partial["record"] + 1},
    )
    book = load_workbook(io.BytesIO(history.content))
    assert sheet_rows(book, "Riwayat Pilihan") == []
    book.close()


def test_xlsx_reads_saved_evidence_and_does_not_attribute_other_coordinates(
    client: TestClient, db_session: Session, partial: dict
) -> None:
    choose(client, partial)
    event = (
        db_session.query(CorrectionDecision)
        .filter_by(process_id=partial["process_id"])
        .one()
    )
    candidate = dict(event.candidate)
    observation = dict(candidate["observations"][0])
    observation["filename"] = "master-berbeda.xlsx"
    observation["source_sheet"] = "Sheet Lain"
    observation["row_number"] = 888
    observation["fields"] = {**observation["fields"], "lat": -7.0}
    candidate["observations"] = [*candidate["observations"], observation]
    event.candidate = candidate
    original = candidate["observations"][0]
    batch = db_session.get(UploadBatch, original["batch_id"])
    batch.filename = "changed-after-choice.xlsx"
    db_session.commit()
    response = client.get(partial["base"] + "/results/export")
    book = load_workbook(io.BytesIO(response.content))
    row = sheet_rows(book, "Hasil Koreksi")[0]
    assert row["sumber file (sheet)"] == "master-pelanggan.xlsx (Master Pelanggan)"
    assert row["sel (baris ke berapa)"] == "3"
    observations = sheet_rows(book, "Observasi Sumber")
    assert len(observations) == 2
    assert observations[0]["koordinat_dipakai"] == "Ya"
    assert observations[1]["koordinat_dipakai"] == "Tidak"
    book.close()
